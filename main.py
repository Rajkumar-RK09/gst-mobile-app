from kivy.app import App
from kivy.uix.boxlayout import BoxLayout
from kivy.uix.scrollview import ScrollView
from kivy.uix.gridlayout import GridLayout
from kivy.uix.label import Label
from kivy.uix.textinput import TextInput
from kivy.uix.button import Button
from kivy.uix.spinner import Spinner
from kivy.uix.popup import Popup
from kivy.uix.widget import Widget
from kivy.graphics import Color, RoundedRectangle
from kivy.metrics import dp
from kivy.core.window import Window
from kivy.utils import platform

from openpyxl import Workbook, load_workbook

from datetime import datetime, timedelta

import os

# ---------------- ANDROID STORAGE ----------------

if platform == "android":
    from android.permissions import request_permissions, Permission
    from android.storage import primary_external_storage_path

    request_permissions([
        Permission.READ_EXTERNAL_STORAGE,
        Permission.WRITE_EXTERNAL_STORAGE
    ])

# ---------------- THEME ----------------

BG_COLOR = (0.06, 0.07, 0.09, 1)

CARD_COLOR = (0.12, 0.13, 0.16, 1)

BLUE_BTN = (0.16, 0.55, 1, 1)
ORANGE_BTN = (1, 0.55, 0.18, 1)
GREEN_BTN = (0.2, 0.8, 0.35, 1)
RED_BTN = (1, 0.28, 0.28, 1)

TEXT_COLOR = (1, 1, 1, 1)

Window.clearcolor = BG_COLOR

GST_RATE = 0.05

# ---------------- STORAGE PATH ----------------

if platform == "android":

    BASE_PATH = primary_external_storage_path()

    SAVE_FOLDER = os.path.join(
        BASE_PATH,
        "Download",
        "GST Billing"
    )

else:

    SAVE_FOLDER = "GST Billing"

os.makedirs(SAVE_FOLDER, exist_ok=True)

# ---------------- ROUNDED BUTTON ----------------


class RoundedButton(Button):

    def __init__(self, bg_color=(1, 1, 1, 1), **kwargs):

        super().__init__(**kwargs)

        self.background_normal = ""
        self.background_down = ""

        self.background_color = (0, 0, 0, 0)

        self.bg_color = bg_color

        with self.canvas.before:
            Color(*self.bg_color)

            self.rect = RoundedRectangle(
                pos=self.pos,
                size=self.size,
                radius=[dp(18)]
            )

        self.bind(pos=self.update_rect)
        self.bind(size=self.update_rect)

    def update_rect(self, *args):

        self.rect.pos = self.pos
        self.rect.size = self.size

# ---------------- GST CALCULATION ----------------


def calculate_gst(amount):

    taxable = amount / (1 + GST_RATE)

    gst = amount - taxable

    cgst = gst / 2
    sgst = gst / 2

    return (
        round(taxable, 2),
        round(cgst, 2),
        round(sgst, 2),
        round(gst, 2)
    )

# ---------------- MAIN APP ----------------


class GSTBillingApp(App):

    def build(self):

        self.current_date = ""

        self.phonepe_bill_no = 1
        self.cash_bill_no = 1

        self.total_sale = 0
        self.total_taxable = 0

        self.file_name = ""

        root = BoxLayout(
            orientation="vertical",
            padding=dp(12),
            spacing=dp(10)
        )

        # ---------------- TITLE ----------------

        title = Label(
            text="GST BILLING SYSTEM",
            font_size="24sp",
            bold=True,
            size_hint_y=None,
            height=dp(60),
            color=TEXT_COLOR
        )

        root.add_widget(title)

        # ---------------- SCROLL ----------------

        scroll = ScrollView()

        layout = GridLayout(
            cols=1,
            spacing=dp(12),
            padding=dp(5),
            size_hint_y=None
        )

        layout.bind(minimum_height=layout.setter('height'))

        scroll.add_widget(layout)

        root.add_widget(scroll)

        # ---------------- DATE ----------------

        self.date_input = TextInput(
            hint_text="Enter Date (DD-MM-YY)",
            multiline=False,
            size_hint_y=None,
            height=dp(58),
            font_size="18sp",
            background_color=CARD_COLOR,
            foreground_color=TEXT_COLOR,
            cursor_color=TEXT_COLOR
        )

        layout.add_widget(self.date_input)

        # ---------------- BILL NUMBERS ----------------

        self.phonepe_input = TextInput(
            hint_text="Starting PhonePe Bill Number",
            multiline=False,
            input_filter="int",
            size_hint_y=None,
            height=dp(58),
            font_size="18sp",
            background_color=CARD_COLOR,
            foreground_color=TEXT_COLOR
        )

        layout.add_widget(self.phonepe_input)

        self.cash_input = TextInput(
            hint_text="Starting Cash Bill Number",
            multiline=False,
            input_filter="int",
            size_hint_y=None,
            height=dp(58),
            font_size="18sp",
            background_color=CARD_COLOR,
            foreground_color=TEXT_COLOR
        )

        layout.add_widget(self.cash_input)

        # ---------------- PAYMENT MODE ----------------

        self.mode_spinner = Spinner(
            text="PhonePe",
            values=("PhonePe", "Cash"),
            size_hint_y=None,
            height=dp(58),
            font_size="18sp",
            background_color=BLUE_BTN
        )

        layout.add_widget(self.mode_spinner)

        # ---------------- SALE AMOUNT ----------------

        self.amount_input = TextInput(
            hint_text="Enter Sale Amount",
            multiline=False,
            input_filter="float",
            size_hint_y=None,
            height=dp(58),
            font_size="18sp",
            background_color=CARD_COLOR,
            foreground_color=TEXT_COLOR
        )

        layout.add_widget(self.amount_input)

        # ---------------- ADD BILL ----------------

        add_btn = RoundedButton(
            text="ADD BILL",
            bg_color=BLUE_BTN,
            size_hint_y=None,
            height=dp(60),
            font_size="20sp",
            bold=True,
            color=TEXT_COLOR
        )

        add_btn.bind(on_press=self.add_bill)

        layout.add_widget(add_btn)

        # ---------------- VIEW & EDIT ----------------

        edit_btn = RoundedButton(
            text="VIEW & EDIT BILLS",
            bg_color=ORANGE_BTN,
            size_hint_y=None,
            height=dp(60),
            font_size="20sp",
            bold=True,
            color=TEXT_COLOR
        )

        edit_btn.bind(on_press=self.show_bills)

        layout.add_widget(edit_btn)

        # ---------------- FINISH DAY ----------------

        finish_btn = RoundedButton(
            text="FINISH DAY",
            bg_color=GREEN_BTN,
            size_hint_y=None,
            height=dp(60),
            font_size="20sp",
            bold=True,
            color=TEXT_COLOR
        )

        finish_btn.bind(on_press=self.finish_day)

        layout.add_widget(finish_btn)

        # ---------------- EXIT ----------------

        exit_btn = RoundedButton(
            text="SAVE & EXIT",
            bg_color=RED_BTN,
            size_hint_y=None,
            height=dp(60),
            font_size="20sp",
            bold=True,
            color=TEXT_COLOR
        )

        exit_btn.bind(on_press=self.exit_app)

        layout.add_widget(exit_btn)

        # ---------------- SUMMARY ----------------

        self.summary_label = Label(
            text="Daily Totals Will Appear Here",
            size_hint_y=None,
            height=dp(220),
            font_size="18sp",
            halign="left",
            valign="top",
            color=TEXT_COLOR
        )

        self.summary_label.bind(
            size=self.summary_label.setter('text_size')
        )

        layout.add_widget(self.summary_label)

        return root

    # ---------------- GET MONTH FILE ----------------

    def create_month_file(self, date_text):

        date_obj = datetime.strptime(date_text, "%d-%m-%y")

        month_name = date_obj.strftime("%B")

        self.file_name = os.path.join(
            SAVE_FOLDER,
            f"GST_Billing_{month_name}.xlsx"
        )

        if not os.path.exists(self.file_name):

            wb = Workbook()

            sheet = wb.active

            sheet.title = "Billing_Data"

            headers = [
                "Date",
                "Payment Mode",
                "Bill No",
                "Sale Amount",
                "Taxable Amount",
                "CGST",
                "SGST",
                "Total GST"
            ]

            sheet.append(headers)

            wb.save(self.file_name)

    # ---------------- POPUP ----------------

    def show_popup(self, title, message):

        box = BoxLayout(
            orientation="vertical",
            spacing=dp(10),
            padding=dp(10)
        )

        msg = Label(
            text=message,
            color=TEXT_COLOR
        )

        btn = RoundedButton(
            text="OK",
            bg_color=BLUE_BTN,
            size_hint_y=None,
            height=dp(52),
            color=TEXT_COLOR
        )

        box.add_widget(msg)
        box.add_widget(btn)

        popup = Popup(
            title=title,
            content=box,
            size_hint=(0.85, 0.45)
        )

        btn.bind(on_press=popup.dismiss)

        popup.open()

    # ---------------- ADD BILL ----------------

    def add_bill(self, instance):

        try:

            if self.current_date == "":

                self.current_date = self.date_input.text.strip()

                datetime.strptime(
                    self.current_date,
                    "%d-%m-%y"
                )

                self.phonepe_bill_no = int(
                    self.phonepe_input.text
                )

                self.cash_bill_no = int(
                    self.cash_input.text
                )

                self.create_month_file(self.current_date)

            amount = float(self.amount_input.text)

            taxable, cgst, sgst, gst = calculate_gst(amount)

            mode = self.mode_spinner.text

            if mode == "PhonePe":

                bill_no = self.phonepe_bill_no

                self.phonepe_bill_no += 1

            else:

                bill_no = self.cash_bill_no

                self.cash_bill_no += 1

            self.total_sale += amount
            self.total_taxable += taxable

            wb = load_workbook(self.file_name)

            sheet = wb.active

            sheet.append([
                self.current_date,
                mode,
                bill_no,
                amount,
                taxable,
                cgst,
                sgst,
                gst
            ])

            wb.save(self.file_name)

            self.summary_label.text = (
                f"Last Bill Added\n\n"
                f"Mode: {mode}\n"
                f"Bill No: {bill_no}\n"
                f"Sale Amount: ₹ {amount:.2f}\n"
                f"Taxable Amount: ₹ {taxable:.2f}\n"
                f"CGST: ₹ {cgst:.2f}\n"
                f"SGST: ₹ {sgst:.2f}\n"
                f"Total GST: ₹ {gst:.2f}\n\n"
                f"DAY SALE TOTAL: ₹ {self.total_sale:.2f}\n"
                f"DAY TAXABLE TOTAL: ₹ {self.total_taxable:.2f}"
            )

            self.amount_input.text = ""

        except Exception as e:

            self.show_popup("Error", str(e))

    # ---------------- FINISH DAY ----------------

    def finish_day(self, instance):

        try:

            wb = load_workbook(self.file_name)

            sheet = wb.active

            sheet.append([""])

            wb.save(self.file_name)

            self.show_popup(
                "Day Finished",
                f"Date: {self.current_date}\n\n"
                f"Total Sale: ₹ {self.total_sale:.2f}\n"
                f"Total Taxable: ₹ {self.total_taxable:.2f}"
            )

            old_date = datetime.strptime(
                self.current_date,
                "%d-%m-%y"
            )

            new_date = old_date + timedelta(days=1)

            self.current_date = new_date.strftime("%d-%m-%y")

            self.date_input.text = self.current_date

            # RESET TO PHONEPE DAILY
            self.mode_spinner.text = "PhonePe"

            self.total_sale = 0
            self.total_taxable = 0

        except Exception as e:

            self.show_popup("Error", str(e))

    # ---------------- SHOW BILLS ----------------

    def show_bills(self, instance):

        try:

            wb = load_workbook(self.file_name)

            sheet = wb.active

            content = BoxLayout(
                orientation="vertical",
                spacing=dp(10),
                padding=dp(10)
            )

            scroll = ScrollView()

            bills_layout = GridLayout(
                cols=1,
                spacing=dp(10),
                size_hint_y=None
            )

            bills_layout.bind(
                minimum_height=bills_layout.setter('height')
            )

            for row in sheet.iter_rows(min_row=2, values_only=True):

                if row[0] is None:
                    continue

                text = (
                    f"{row[0]} | "
                    f"{row[1]} | "
                    f"Bill {row[2]} | "
                    f"₹ {row[3]}"
                )

                lbl = Label(
                    text=text,
                    size_hint_y=None,
                    height=dp(42),
                    font_size="16sp",
                    color=TEXT_COLOR,
                    halign="left",
                    valign="middle"
                )

                lbl.bind(size=lbl.setter('text_size'))

                bills_layout.add_widget(lbl)

            scroll.add_widget(bills_layout)

            content.add_widget(scroll)

            mode_spinner = Spinner(
                text="PhonePe",
                values=("PhonePe", "Cash"),
                size_hint_y=None,
                height=dp(52)
            )

            bill_input = TextInput(
                hint_text="Enter Bill Number",
                multiline=False,
                input_filter="int",
                size_hint_y=None,
                height=dp(52)
            )

            amount_input = TextInput(
                hint_text="New Sale Amount",
                multiline=False,
                input_filter="float",
                size_hint_y=None,
                height=dp(52)
            )

            save_btn = RoundedButton(
                text="SAVE CHANGES",
                bg_color=GREEN_BTN,
                size_hint_y=None,
                height=dp(55),
                color=TEXT_COLOR
            )

            content.add_widget(mode_spinner)
            content.add_widget(bill_input)
            content.add_widget(amount_input)
            content.add_widget(save_btn)

            popup = Popup(
                title="View & Edit Bills",
                content=content,
                size_hint=(0.95, 0.95)
            )

            def save_changes(btn):

                try:

                    mode = mode_spinner.text

                    bill_no = int(bill_input.text)

                    new_amount = float(amount_input.text)

                    taxable, cgst, sgst, gst = calculate_gst(
                        new_amount
                    )

                    found = False

                    for row in sheet.iter_rows(min_row=2):

                        if (
                            str(row[1].value).lower()
                            == mode.lower()
                            and row[2].value == bill_no
                        ):

                            row[3].value = new_amount
                            row[4].value = taxable
                            row[5].value = cgst
                            row[6].value = sgst
                            row[7].value = gst

                            found = True

                            break

                    wb.save(self.file_name)

                    if found:

                        self.show_popup(
                            "Success",
                            "Bill updated successfully."
                        )

                    else:

                        self.show_popup(
                            "Error",
                            "Bill not found."
                        )

                    popup.dismiss()

                except Exception as e:

                    self.show_popup("Error", str(e))

            save_btn.bind(on_press=save_changes)

            popup.open()

        except Exception as e:

            self.show_popup("Error", str(e))

    # ---------------- EXIT APP ----------------

    def exit_app(self, instance):

        self.show_popup(
            "Saved Successfully",
            f"Excel File Saved\n\n"
            f"Location:\n{self.file_name}"
        )

        App.get_running_app().stop()


GSTBillingApp().run()